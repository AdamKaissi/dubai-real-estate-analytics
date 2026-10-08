"""Step 5: export small summary tables from the database for the Power BI web service.
Writes outputs/bi/dubai_re_bi.xlsx (4 sheets, each a real Excel table) plus CSV copies.
Run 04b_stats_v2.py first so the regression table can be included."""
import os
import numpy as np
import pandas as pd
from sqlalchemy import create_engine
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from config import DATABASE_URL

os.makedirs("outputs/bi", exist_ok=True)
engine = create_engine(DATABASE_URL)

df = pd.read_sql("""
    SELECT d.full_date, d.year, d.quarter, a.area_name, p.reg_type, f.project_name,
           f.value_aed, f.price_per_sqm
    FROM fact_transactions f
    JOIN dim_date d ON d.date_id = f.date_id
    JOIN dim_area a ON a.area_id = f.area_id
    JOIN dim_property p ON p.property_id = f.property_id
    WHERE NOT f.is_bulk""", engine, parse_dates=["full_date"])
print(f"Read {len(df):,} non-bulk sales")

# Only mappings I am confident about. Add your own after checking sheet 'area_projects'.
LABELS = {"Marsa Dubai": "Marsa Dubai (Dubai Marina)", "Burj Khalifa": "Burj Khalifa (Downtown Dubai)"}
df["area_label"] = df["area_name"].replace(LABELS)
df["reg_type"] = df["reg_type"].replace({"Existing Properties": "Existing", "Off-Plan Properties": "Off-plan"})

# 1. Quarterly market summary (drop an incomplete final quarter)
df["period"] = df["year"].astype(str) + " Q" + df["quarter"].astype(str)
last = df["full_date"].max()
q_end = pd.Period(last, "Q").end_time.normalize()
if last.normalize() < q_end:
    cutoff = pd.Period(last, "Q")
    df_q = df[df["full_date"].dt.to_period("Q") < cutoff]
    print(f"Excluded incomplete quarter {cutoff} (data ends {last.date()})")
else:
    df_q = df
quarterly = (df_q.groupby(["year", "quarter", "period", "reg_type"])
             .agg(n_sales=("price_per_sqm", "size"), median_psqm=("price_per_sqm", "median"),
                  total_value_aed=("value_aed", "sum")).reset_index())
quarterly[["median_psqm", "total_value_aed"]] = quarterly[["median_psqm", "total_value_aed"]].round(0)

# 2. Area x year x registration type (cells with >= 30 sales)
area_year = (df.groupby(["area_label", "year", "reg_type"])
             .agg(n_sales=("price_per_sqm", "size"), median_psqm=("price_per_sqm", "median")).reset_index())
area_year = area_year[area_year["n_sales"] >= 30].copy()
area_year["median_psqm"] = area_year["median_psqm"].round(0)

# 3. Regression effects, converted to plain percentages
reg_path = "outputs/regression_v2_coefficients.csv"
rows = []
if os.path.exists(reg_path):
    r = pd.read_csv(reg_path, index_col=0)
    ci_lo, ci_hi = np.log1p(r["ci_low_pct"] / 100), np.log1p(r["ci_high_pct"] / 100)
    for term, row in r.iterrows():
        if "year_c" in term:
            g, lab, k = "Year vs 2019", term.split("T.")[1].rstrip("]"), 1
        elif "area_grp" in term and "Other" not in term:
            g, lab, k = "Area vs baseline area", term.split("T.")[1].rstrip("]"), 1
        elif "reg_type" in term:
            g, lab, k = "Registration type", "Off-plan vs existing", 1
        elif "property_sub_type" in term:
            g, lab, k = "Property type", "Villa vs flat", 1
        elif term == "log_area":
            g, lab, k = "Unit size", "Doubling unit size", np.log(2)   # coef is per log-unit
        else:
            continue
        beta, lo, hi = np.log1p(row["pct_effect"] / 100), ci_lo[term], ci_hi[term]
        rows.append({"group": g, "label": lab,
                     "pct_effect": round((np.exp(beta * k) - 1) * 100, 1),
                     "ci_low_pct": round((np.exp(lo * k) - 1) * 100, 1),
                     "ci_high_pct": round((np.exp(hi * k) - 1) * 100, 1)})
else:
    print("NOTE: outputs/regression_v2_coefficients.csv not found; run 04b_stats_v2.py first.")
effects = pd.DataFrame(rows, columns=["group", "label", "pct_effect", "ci_low_pct", "ci_high_pct"])

# 4. Top 3 projects in each of the 25 biggest areas (use this to label areas honestly)
top_areas = df["area_name"].value_counts().head(25).index
ap = (df[df["area_name"].isin(top_areas)].groupby(["area_name", "project_name"]).size()
      .reset_index(name="n_sales").sort_values(["area_name", "n_sales"], ascending=[True, False]))
ap = ap.groupby("area_name").head(3)

sheets = {"quarterly_market": quarterly, "area_year": area_year,
          "regression_effects": effects, "area_projects": ap}
wb = Workbook(); wb.remove(wb.active)
for name, t in sheets.items():
    t.to_csv(f"outputs/bi/{name}.csv", index=False)
    ws = wb.create_sheet(name)
    ws.append(list(t.columns))
    for rec in t.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v) for v in rec])
    if len(t):
        ref = f"A1:{get_column_letter(len(t.columns))}{len(t) + 1}"
        tab = Table(displayName=name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    print(f"{name:20s} {len(t):>7,} rows")
wb.save("outputs/bi/dubai_re_bi.xlsx")
print("Saved outputs/bi/dubai_re_bi.xlsx")
